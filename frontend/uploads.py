"""对话附件：从上传文件中提取纯文本，供 AI 作为上下文参考。
支持：纯文本类、PDF、Word(.docx)、Excel(.xlsx/.csv)。缺少对应库时该类型优雅降级。"""
import csv
import io
import os

MAX_CHARS_PER_FILE = 20000   # 单文件提取上限
MAX_FILES = 5

TEXT_EXTS = {".txt", ".md", ".markdown", ".csv", ".tsv", ".log", ".json", ".ini", ".cfg",
             ".conf", ".yaml", ".yml", ".xml", ".html", ".htm", ".c", ".h", ".cpp", ".hpp",
             ".py", ".js", ".java", ".sql", ".sh", ".bat", ".v", ".sv", ".vhd"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
XLSX_EXTS = {".xlsx", ".xlsm"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}

ALLOWED_EXTS = TEXT_EXTS | PDF_EXTS | DOCX_EXTS | XLSX_EXTS | IMAGE_EXTS


def _clip(text):
    text = (text or "").strip()
    if len(text) > MAX_CHARS_PER_FILE:
        return text[:MAX_CHARS_PER_FILE], True
    return text, False


def _from_pdf(data):
    import fitz  # PyMuPDF
    parts = []
    with fitz.open(stream=data, filetype="pdf") as doc:
        for i, page in enumerate(doc):
            parts.append(page.get_text())
            if sum(len(p) for p in parts) > MAX_CHARS_PER_FILE:
                break
    return "\n".join(parts)


def _from_docx(data):
    import docx  # python-docx
    d = docx.Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _from_xlsx(data):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"# 工作表: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            if any(cells):
                out.append(" | ".join(cells))
            if sum(len(x) for x in out) > MAX_CHARS_PER_FILE:
                break
    return "\n".join(out)


def extract_text(filename, data):
    """返回 (text, truncated)。无法解析时抛出 ValueError。"""
    ext = os.path.splitext(filename or "")[1].lower()
    if ext in PDF_EXTS:
        try:
            return _clip(_from_pdf(data))
        except ImportError:
            raise ValueError("服务器未安装 PDF 解析库(PyMuPDF)")
    if ext in DOCX_EXTS:
        try:
            return _clip(_from_docx(data))
        except ImportError:
            raise ValueError("服务器未安装 Word 解析库(python-docx)")
    if ext in XLSX_EXTS:
        try:
            return _clip(_from_xlsx(data))
        except ImportError:
            raise ValueError("服务器未安装 Excel 解析库(openpyxl)")
    if ext in TEXT_EXTS or ext == "":
        if ext == ".csv":
            try:
                rows = list(csv.reader(io.StringIO(data.decode("utf-8", "ignore"))))
                return _clip("\n".join(" | ".join(r) for r in rows))
            except Exception:
                pass
        return _clip(data.decode("utf-8", "ignore"))
    raise ValueError(f"不支持的文件类型: {ext or '未知'}")


ATTACH_HEADER = "〔以下为用户上传的附件，供分析参考〕"


def build_message(message, attachments):
    """attachments: list of (filename, text, truncated)。拼成发给后端的消息。"""
    message = (message or "").strip()
    if not attachments:
        return message
    parts = [message] if message else []
    parts.append(ATTACH_HEADER)
    for name, text, truncated in attachments:
        body = text if text else "(未能提取到文本内容)"
        if truncated:
            body += "\n…(内容过长，已截断)"
        parts.append(f"【文件：{name}】\n{body}")
    return "\n".join(parts)
